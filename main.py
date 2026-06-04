# main.py
import os
import random
import telebot
from telebot import types
from dotenv import load_dotenv
from google import genai
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from gtts import gTTS

# Borkunii tokenho
load_dotenv()
BOT_TOKEN = os.getenv("BOT_TOKEN")
GEMINI_KEY = os.getenv("GEMINI_API_KEY")

client = genai.Client(api_key=GEMINI_KEY)
MODEL_NAME = 'gemini-2.5-flash'

bot = telebot.TeleBot(BOT_TOKEN)

try:
    from science_data import SCIENCE_DATABASE, DICTIONARY
except ImportError:
    SCIENCE_DATABASE = {}
    DICTIONARY = {}

user_states = {}
user_scores = {}

def create_ariza_docx(filename, to_whom, from_whom, title, body, date_sign):
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(0.8)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)

    p_header1 = doc.add_paragraph()
    p_header1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_header1.paragraph_format.left_indent = Inches(3.0)
    p_header1.add_run(to_whom + "\n")
    
    p_header2 = doc.add_paragraph()
    p_header2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_header2.paragraph_format.left_indent = Inches(3.0)
    p_header2.add_run(from_whom + "\n\n")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(title)
    run_title.bold = True
    run_title.font.size = Pt(16)
    doc.add_paragraph()

    p_body = doc.add_paragraph()
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body.paragraph_format.first_line_indent = Inches(0.5)
    p_body.add_run(body + "\n\n")

    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_footer.add_run(date_sign)
    doc.save(filename)

@bot.message_handler(commands=['start', 'menu'])
def send_welcome(message):
    user_states[message.chat.id] = {"path": [], "mode": "science", "step_flow": "alphabet", "level": 1}
    show_main_menu(message.chat.id)

def show_main_menu(chat_id):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    markup.add(types.KeyboardButton("🔬 Бахшҳои Илмӣ"), types.KeyboardButton("🕌 Хати Ниёгон (Форсӣ)"))
    markup.add(types.KeyboardButton("📄 Санаднавис ва Ариза"), types.KeyboardButton("📖 Фарҳанг ва Луғат"))
    
    bot.send_message(
        chat_id, 
        "🧠 **Ассистенти Мукаммали AI ва Илмӣ Омода Аст!**\n\nЛутфан яке аз бахшҳои муосири зеринро интихоб кунед:", 
        parse_mode="Markdown",
        reply_markup=markup
    )

def show_science_sub_menu(chat_id):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    markup.add(types.KeyboardButton("📐 Масъалаи Математика"), types.KeyboardButton("🧪 Реаксияи Химия"))
    markup.add(types.KeyboardButton("🌌 Масъалаи Физика"), types.KeyboardButton("📝 Таҳлили Шеър ва Ҷумла"))
    markup.add(types.KeyboardButton("🇺🇸 Тоҷикӣ ➡️ Англисӣ"), types.KeyboardButton("🇷🇺 Тоҷикӣ ➡️ Русӣ"))
    markup.add(types.KeyboardButton("🇹🇯 Русӣ/Англисӣ ➡️ Тоҷикӣ"))
    markup.add(types.KeyboardButton("🇬🇧 Омӯзиши Англисӣ"), types.KeyboardButton("🇷🇺 Омӯзиши Русӣ"))
    markup.add(types.KeyboardButton("🇹🇯 Омӯзиши Тоҷикӣ"))
    markup.add(types.KeyboardButton("⬅️ Ба Менюи Асосӣ"))
    bot.send_message(chat_id, "🤖 **Бахши Илм ва Забонҳо.** Кадом намуди кӯмак лозим аст?", parse_mode="Markdown", reply_markup=markup)

def show_persian_menu(chat_id):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=1)
    markup.add(types.KeyboardButton("🔄 Баргардонии Матн (Кириллӣ ➡️ Форсӣ)"))
    markup.add(types.KeyboardButton("🎮 Омӯзиши Хати Форсӣ (Duolingo)"))
    markup.add(types.KeyboardButton("⬅️ Ба Менюи Асосӣ"))
    bot.send_message(chat_id, "🕌 **Бахши Хати Ниёгон.** Навъи корро интихоб кунед:", parse_mode="Markdown", reply_markup=markup)

def show_ariza_menu(chat_id):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=1)
    markup.add(types.KeyboardButton("📝 Навиштани Аризаи Нав (AI)"))
    markup.add(types.KeyboardButton("⬅️ Ба Менюи Асосӣ"))
    bot.send_message(chat_id, "📄 **Ёрдамчии Коргузорӣ.** Барои сохтани ариза дар формати Word тугмаро пахш кунед:", parse_mode="Markdown", reply_markup=markup)


# ================= ЗАНҶИРИ НАВИ ОМӮЗИШӢ (АЛИФБО -> КАЛИМА -> ТЕСТИ КОМБИНИРОНИДАШУДА) =================
def send_dynamic_duolingo_step(chat_id):
    if chat_id not in user_states:
        user_states[chat_id] = {"path": [], "mode": "duolingo_fa", "step_flow": "alphabet", "level": 1}
        
    state = user_states[chat_id]
    current_flow = state.get("step_flow", "alphabet")
    level = state.get("level", 1)

    # 1. ҚАДАМИ АЛИФБО (ҲАРФҲО + ОВОЗ)
    if current_flow == "alphabet":
        bot.send_message(chat_id, f"⏳ Устоди AI дарси ҳарфҳои алифборо (Сатҳи {level}) омода мекунад...")
        prompt = (
            f"Ту муаллими касбии хати форсӣ ҳастӣ. Як ҳарфи алифборо вобаста ба Сатҳи {level} интихоб кун, "
            f"тарзи навишт, ном ва садояшро ба тоҷикии кириллӣ фаҳмон. Дар охир ХАТМАН танҳо номи худи ҳарфро "
            f"дар чунин формат гузор: [AUDIO] номи ҳарф ё садояш [AUDIO]. Масалан: [AUDIO] الف [AUDIO] ё [AUDIO] ب [AUDIO]"
        )
        try:
            response = client.models.generate_content(model=MODEL_NAME, contents=prompt)
            res_text = response.text
            
            text_to_speak = "الف"
            if "[AUDIO]" in res_text:
                parts = res_text.split("[AUDIO]")
                res_text = parts[0] + (parts[2] if len(parts) > 2 else "")
                text_to_speak = parts[1].strip()

            state["step_flow"] = "word" # Қадами навбатиро ба калима мегузаронем
            markup = types.ReplyKeyboardMarkup(resize_keyboard=True).add("⏭️ Гузаштан ба калимаҳо", "⬅️ Ба Менюи Асосӣ")
            bot.send_message(chat_id, f"📖 **ҚАДАМИ 1: ОМӮЗИШИ ҲАРФИ АЛИФБО**\n\n{res_text}", reply_markup=markup)

            # Овоздиҳии ҳарф
            tts = gTTS(text=text_to_speak, lang='fa')
            audio_file = f"alphabet_{chat_id}.mp3"
            tts.save(audio_file)
            with open(audio_file, 'rb') as audio:
                bot.send_voice(chat_id, audio, caption=f"🎧 Талаффузи ҳарф: {text_to_speak}")
            os.remove(audio_file)
        except Exception as e:
            bot.send_message(chat_id, f"❌ Хатогӣ: {e}")
            state["step_flow"] = "alphabet"

    # 2. ҚАДАМИ КАЛИМАҲО (МАТН + ТАЛАФФУЗИ БОТ)
    elif current_flow == "word":
        bot.send_message(chat_id, f"⏳ Устоди AI дарс ва калимаҳои навсохтаро омода мекунад...")
        prompt = (
            f"Дар Сатҳи {level} якчанд калимаҳои соддаи форсиро, ки бо ҳарфҳои омӯхташуда сохта мешаванд, бо тарҷумаи кириллиашон нишон деҳ. "
            f"Дар охир як калимаи асосиро барои овоздиҳӣ дар формат равон кун: [AUDIO] калимаи форсӣ [AUDIO]."
        )
        try:
            response = client.models.generate_content(model=MODEL_NAME, contents=prompt)
            res_text = response.text
            
            text_to_speak = "آب"
            if "[AUDIO]" in res_text:
                parts = res_text.split("[AUDIO]")
                res_text = parts[0] + (parts[2] if len(parts) > 2 else "")
                text_to_speak = parts[1].strip()

            state["step_flow"] = "test_combined" # Гузариш ба тести маҷмӯӣ
            markup = types.ReplyKeyboardMarkup(resize_keyboard=True).add("⏭️ Гузаштан ба Имтиҳон (Тест)", "⬅️ Ба Менюи Асосӣ")
            bot.send_message(chat_id, f"📝 **ҚАДАМИ 2: ОМӮЗИШИ КАЛИМАҲО**\n\n{res_text}", reply_markup=markup)

            # Овоздиҳии калима
            tts = gTTS(text=text_to_speak, lang='fa')
            audio_file = f"word_{chat_id}.mp3"
            tts.save(audio_file)
            with open(audio_file, 'rb') as audio:
                bot.send_voice(chat_id, audio, caption=f"🎧 Бот калимаро талаффуз мекунад: {text_to_speak}")
            os.remove(audio_file)
        except Exception as e:
            bot.send_message(chat_id, f"❌ Хатогӣ: {e}")
            state["step_flow"] = "word"

    # 3. ҚАДАМИ ТЕСТИ МАҶМӮӢ (НИШОН ДОДАНИ КАЛИМА + НАВИШТАНИ КИРИЛЛӢ + СУПОРИДАНИ ОВОЗ)
    elif current_flow == "test_combined":
        prompt = f"Як калимаи форсиро барои Сатҳи {level} интихоб кун. ТАНҲО худи калимаи форсиро равон кун (Масалан: 'نان' ё 'مادر'). Ҳеҷ чизи дигар нанавис."
        try:
            response = client.models.generate_content(model=MODEL_NAME, contents=prompt)
            target_word = response.text.strip()
            
            state["test_word_target"] = target_word
            state["waiting_for_test_answer"] = True
            
            markup = types.ReplyKeyboardMarkup(resize_keyboard=True).add("⏭️ Гузаштани ин тест", "⬅️ Ба Менюи Асосӣ")
            bot.send_message(
                chat_id, 
                f"🎯 **ҚАДАМИ 3: ИМТИҲОНИ МАҶМӮӢ (ПРОФЕССИОНАЛӢ)**\n\n"
                f"Калимаи форсиро бинед: 👉 **{target_word}**\n\n"
                f"**Вазифаи шумо:**\n"
                f"1. Маъно ё шакли кириллии ин калимаро бо матни оддӣ навишта фиристед.\n"
                f"2. Сипас, тугмаи микрофонро зер карда, овози худро (Voice) ҳангоми хондани ин калима ба бот равон кунед!", 
                parse_mode="Markdown", 
                reply_markup=markup
            )
        except:
            state["step_flow"] = "alphabet"
            send_dynamic_duolingo_step(chat_id)


# ================= KORKARDI МАТНҲО ВА ТУГМАҲО =================
@bot.message_handler(func=lambda message: True)
def handle_all_messages(message):
    chat_id = message.chat.id
    text = message.text
    
    if chat_id not in user_states:
        user_states[chat_id] = {"path": [], "mode": "science", "step_flow": "alphabet", "level": 1}
    if chat_id not in user_scores:
        user_scores[chat_id] = 0
        
    state = user_states[chat_id]

    if text in ["⬅️ Ба Менюи Асосӣ", "⬅️ Бозгашт ба Меню", "⬅️ Бозгашт"]:
        state["mode"] = "science"
        show_main_menu(chat_id)
        return

    if text == "🔬 Бахшҳои Илмӣ":
        show_science_sub_menu(chat_id)
        return
        
    if text == "🕌 Хати Ниёгон (Форсӣ)":
        show_persian_menu(chat_id)
        return
        
    if text == "📄 Санаднавис ва Ариза":
        show_ariza_menu(chat_id)
        return

    if text == "📖 Фарҳанг ва Луғат":
        state["mode"] = "dictionary"
        markup = types.ReplyKeyboardMarkup(resize_keyboard=True).add("⬅️ Ба Менюи Асосӣ")
        bot.send_message(chat_id, "🔍 **Калимаи лозимиро бинавесед:**", parse_mode="Markdown", reply_markup=markup)
        return

    if text == "🎮 Омӯзиши Хати Форсӣ (Duolingo)":
        state["mode"] = "duolingo_fa"
        state["step_flow"] = "alphabet"  # Оғоз аз тартиби алифбо
        bot.send_message(chat_id, f"🎮 **Хуш омадед ба мактаби касбии хати форсӣ!**\n🌟 Холҳои шумо: {user_scores[chat_id]}")
        send_dynamic_duolingo_step(chat_id)
        return

    # МАНТИҚИ ТЕСТИ МАТНИИ ДУОЛИНГО
    if state["mode"] == "duolingo_fa":
        if text in ["⏭️ Гузаштан ба калимаҳо", "⏭️ Гузаштан ба Имтиҳон (Тест)", "⏭️ Дарси Навбатӣ", "⏭️ Гузаштани ин тест"]:
            send_dynamic_duolingo_step(chat_id)
            return
            
        if state.get("waiting_for_test_answer") and "test_word_target" in state:
            target = state["test_word_target"]
            bot.send_message(chat_id, "⏳ Тексти фиристодаи шумо қабул шуд. Акнун лутфан **Овози худро (Voice)** низ бифристед, то устоди AI онро таҳлил кунад.")
            
            # Санҷиши тарҷумаи матн бо AI
            prompt = f"Корбар барои калимаи форсии '{target}' тарҷумаи кириллии '{text}'-ро навишт. Оё ин тарҷума дуруст аст? Танҳо бо 'Бале' ё 'Не' ҷавоб деҳ ва шарҳи кӯтоҳ навис."
            try:
                res = client.models.generate_content(model=MODEL_NAME, contents=prompt)
                bot.send_message(chat_id, f"📝 **Натиҷаи матнии тест:**\n{res.text}")
                if "бале" in res.text.lower():
                    user_scores[chat_id] += 10
            except:
                pass
            return

    # БАХШҲОИ САНАДНАВИСӢ ВА ИЛМҲО (НЕСТ КАРДА НАШУДААСТ)
    if text == "📝 Навиштани Аризаи Нав (AI)":
        state["mode"] = "make_ariza"
        markup = types.ReplyKeyboardMarkup(resize_keyboard=True).add("⬅️ Ба Менюи Асосӣ")
        bot.send_message(chat_id, "📄 **Шарти аризаро нависед:**", parse_mode="Markdown", reply_markup=markup)
        return

    if text == "🔄 Баргардонии Матн (Кириллӣ ➡️ Форсӣ)":
        state["mode"] = "translit_fa"
        markup = types.ReplyKeyboardMarkup(resize_keyboard=True).add("⬅️ Ба Менюи Асосӣ")
        bot.send_message(chat_id, "🔄 **Матнро бо хати кириллӣ равон кунед:**", parse_mode="Markdown", reply_markup=markup)
        return

    if text in ["📝 Таҳлили Шеър ва Ҷумла", "📐 Масъалаи Математика", "🌌 Масъалаи Физика", "🧪 Реаксияи Химия", "🇺🇸 Тоҷикӣ ➡️ Англисӣ", "🇷🇺 Тоҷикӣ ➡️ Русӣ", "🇹🇯 Русӣ/Англисӣ ➡️ Тоҷикӣ", "🇬🇧 Омӯзиши Англисӣ", "🇷🇺 Омӯзиши Русӣ", "🇹🇯 Омӯзиши Тоҷикӣ"]:
        ai_map = {
            "📝 Таҳлили Шеър ва Ҷумла": "ai_literature", "📐 Масъалаи Математика": "ai_math",
            "🌌 Масъалаи Физика": "ai_physics", "🧪 Реаксияи Химия": "ai_chemistry",
            "🇺🇸 Тоҷикӣ ➡️ Англисӣ": "tr_tj_en", "🇷🇺 Тоҷикӣ ➡️ Русӣ": "tr_tj_ru",
            "🇹🇯 Русӣ/Англисӣ ➡️ Тоҷикӣ": "tr_any_tj", "🇬🇧 Омӯзиши Англисӣ": "lang_en",
            "🇷🇺 Омӯзиши Русӣ": "lang_ru", "🇹🇯 Омӯзиши Тоҷикӣ": "lang_tj"
        }
        state["mode"] = ai_map[text]
        markup = types.ReplyKeyboardMarkup(resize_keyboard=True).add("⬅️ Ба Менюи Асосӣ")
        bot.send_message(chat_id, f"🚀 Бахши **{text}** фаъол шуд. Саволро бинавесед:", parse_mode="Markdown", reply_markup=markup)
        return

    if state["mode"] == "make_ariza":
        bot.send_message(chat_id, "⏳ Ариза сохта шуда истодааст...")
        prompt = f"Ту ассистенти ҳуқуқӣ ҳастӣ. Барои ин дархост аризаи расмии тоҷикӣ соз: {text}. Ҷавобро ба 5 қисм бо [SPLIT] ҷудо кун: Ба кӣ[SPLIT]Аз кӣ[SPLIT]АРИЗА[SPLIT]Матн[SPLIT]Сана"
        try:
            response = client.models.generate_content(model=MODEL_NAME, contents=prompt)
            parts = response.text.split("[SPLIT]")
            if len(parts) >= 5:
                filename = f"Ariza_{chat_id}.docx"
                create_ariza_docx(filename, parts[0].strip(), parts[1].strip(), parts[2].strip(), parts[3].strip(), parts[4].strip())
                with open(filename, 'rb') as doc_file:
                    bot.send_document(chat_id, doc_file, caption="📄 Ариза дар формати Word омода шуд!")
                os.remove(filename)
            else: bot.send_message(chat_id, "⚠️ Хатогии формат:\n" + response.text)
        except Exception as e: bot.send_message(chat_id, f"❌ Хатогӣ: {e}")
        return

    if state["mode"] in ["translit_fa", "ai_literature", "ai_math", "ai_physics", "ai_chemistry", "tr_tj_en", "tr_tj_ru", "tr_any_tj", "lang_en", "lang_ru", "lang_tj"]:
        latex_danger_rules = " ХЕЛЕ МУҲИМ: Ҳеҷ гоҳ аз аломатҳои кодҳои LaTeX истифода набар!"
        if state["mode"] == "translit_fa": prompt = f"Ин матни тоҷикиро ба хати форсӣ баргардон: {text}"
        elif state["mode"] == "ai_math": prompt = f"Шарти масъала: {text}. {latex_danger_rules}"
        elif state["mode"] == "ai_physics": prompt = f"Масъалаи физика: {text}. {latex_danger_rules}"
        elif state["mode"] == "ai_chemistry": prompt = f"Реаксия ё масъалаи химия: {text}. {latex_danger_rules}"
        else: prompt = f"Дархост: {text}"
        try:
            response = client.models.generate_content(model=MODEL_NAME, contents=prompt)
            bot.send_message(chat_id, response.text)
        except Exception as e: bot.send_message(chat_id, f"❌ Хатогӣ: {e}")
        return


# ================= ТАҲЛИЛИ СУПОРИШҲОИ ОВОЗӢ СУПОРИДАШУДА (VOICE ANALYSIS) =================
@bot.message_handler(content_types=['voice'])
def handle_voice_message(message):
    chat_id = message.chat.id
    state = user_states.get(chat_id, {})
    
    if state.get("mode") == "duolingo_fa" and "test_word_target" in state:
        bot.send_message(chat_id, "📥 Овози шумо гирифта шуд. Профессори AI онро гӯш карда, талаффузатонро таҳлил дорад...")
        try:
            file_info = bot.get_file(message.voice.file_id)
            downloaded_file = bot.download_file(file_info.file_path)
            
            user_voice_file = f"user_voice_{chat_id}.ogg"
            with open(user_voice_file, 'wb') as new_file:
                new_file.write(downloaded_file)
            
            with open(user_voice_file, 'rb') as f:
                audio_data = f.read()
                
            target = state["test_word_target"]
            prompt = (
                f"Ту профессори забони форсӣ ҳастӣ. Ин файли аудиоии корбар аст, ки мехоҳад калимаи '{target}'-ро хонад. "
                f"Аудиоро гӯш кун ва таҳлили касбӣ кун: "
                f"1. Оё дуруст хонд ва лаҳҷааш соз аст? (Бале ё Не) "
                f"2. Чӣ хатоии фонетикӣ дорад ва чӣ тавр бояд ислоҳ кунад? "
                f"Ҷавобро танҳо бо тоҷикии кириллӣ навишта, дар аввал агар дуруст бошад [CORRECT] ва агар хато бошад [WRONG] гузор."
            )
            
            response = client.models.generate_content(
                model=MODEL_NAME,
                contents=[
                    {"inline_data": {"mime_type": "audio/ogg", "data": audio_data}},
                    prompt
                ]
            )
            
            analysis_result = response.text
            os.remove(user_voice_file)
            
            if "[CORRECT]" in analysis_result:
                user_scores[chat_id] += 20
                bot.send_message(chat_id, f"🎉 **АЪЛО! ТАЛАФФУЗИ ОВОЗИИ ШУМО ДУРУСТ АСТ!**\n+20 хол! 🌟\n\n{analysis_result.replace('[CORRECT]', '')}")
            else:
                bot.send_message(chat_id, f"⚠️ **ТАҲЛИЛИ ТАЛАФФУЗИ ОВОЗӢ АЗ AI:**\n\n{analysis_result.replace('[WRONG]', '')}")
            
            # Навсозии сатҳ ва гузаштан ба давраи нави занҷир (боз аз Алифбо)
            state["level"] = (user_scores[chat_id] // 60) + 1
            state["step_flow"] = "alphabet"
            
            del state["test_word_target"]
            if "waiting_for_test_answer" in state:
                del state["waiting_for_test_answer"]
                
            bot.send_message(chat_id, f"📊 Холҳои умумии шумо: {user_scores[chat_id]}. Ба дарси навбатӣ мегузарем...")
            send_dynamic_duolingo_step(chat_id)
            
        except Exception as e:
            bot.send_message(chat_id, f"❌ Хатогӣ ҳангоми таҳлили овоз: {e}")
    else:
        bot.send_message(chat_id, "🤖 Овози шумо қабул шуд, вале ҳозир дар қисми имтиҳони овозӣ нестед.")

if __name__ == '__main__':
    print("Бот бо занҷири нави дарсӣ (Алифбо -> Калима -> Имтиҳони маҷмӯӣ) онлайн аст!")
    import time
    while True:
        try:
            bot.polling(none_stop=True, timeout=20, long_polling_timeout=20)
        except Exception as e:
            time.sleep(5)
